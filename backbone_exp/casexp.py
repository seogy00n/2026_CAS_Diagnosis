# -*- coding: utf-8 -*-
"""
CAS 항목별 분류 - 백본 비교 실험 공통 모듈.

설계 원칙
  1. 백본은 항상 동결한다. 특징을 한 번 뽑아 캐시하고 헤드만 재학습한다.
  2. 분할은 환자 단위로 고정한다. 같은 시드의 같은 fold를 모든 백본이 공유한다.
  3. 임계값은 학습 fold 안에서만 고른다. 검증 fold는 평가에만 쓴다.
"""
from __future__ import annotations
import json, hashlib
from pathlib import Path

import numpy as np
import torch
import torch.nn as nn
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
IMG_DIR = ROOT / "실습" / "data_origin"
TED_DIR = ROOT / "TED_Data"
CACHE = Path(__file__).resolve().parent / "cache"
CACHE.mkdir(exist_ok=True)

CAS_ITEMS = ["CAS_3", "CAS_4", "CAS_5", "CAS_7"]       # CAS_6은 비정상 0건이라 제외
CAS_NAMES = {"CAS_3": "안검 발적", "CAS_4": "결막 발적",
             "CAS_5": "안검 부종", "CAS_7": "눈물언덕 염증"}
SEED = 20260914


# ---------------------------------------------------------------- 데이터
def load_labels_ted() -> dict[str, dict[str, int]]:
    """TED_Data 폴더 구조에서 환자 단위 항목 라벨을 복원한다.

    <환자>/<CAS_n>/{normal|abnormal}/{left|right}/ 구조이며,
    한 환자 안에서 normal/abnormal이 섞인 항목은 0건임을 사전 검증했다.
    """
    out = {}
    for pdir in sorted(p for p in TED_DIR.iterdir() if p.is_dir()):
        row = {}
        for cas in CAS_ITEMS:
            cdir = pdir / cas
            if not cdir.is_dir():
                raise FileNotFoundError(f"{cdir} 없음")
            row[cas] = 1 if (cdir / "abnormal").is_dir() else 0
        out[pdir.name] = row
    return out


def load_labels_union() -> dict[str, dict[str, int]]:
    """정답지 xlsx의 Union 시트에서 라벨을 읽는다 (20명분만 존재)."""
    import openpyxl
    wb = openpyxl.load_workbook(ROOT / "CAS_Union_Ground_Truth.xlsx", data_only=True)
    ws = wb["종합 정답지(Union)"]
    by_idx = {}
    for r in range(4, ws.max_row + 1):
        rid = ws.cell(r, 1).value
        if not isinstance(rid, str) or not rid.startswith("AI-01-"):
            continue
        idx = int(rid.split("-")[-1])
        vals = [ws.cell(r, c).value for c in range(6, 11)]        # CAS 3,4,5,6,7
        by_idx[idx] = {"CAS_3": vals[0], "CAS_4": vals[1],
                       "CAS_5": vals[2], "CAS_7": vals[4]}
    # AI-01-0NN <-> data_origin 폴더의 N번째(정렬 순서)가 대응한다.
    folders = sorted(p.name for p in IMG_DIR.iterdir() if p.is_dir())
    return {name: by_idx[i] for i, name in enumerate(folders, 1) if i in by_idx}


def list_images(patients: list[str]) -> tuple[list[Path], list[str]]:
    """환자별 좌/우 폴더의 영상 경로와 소속 환자를 반환한다."""
    paths, owners = [], []
    for pid in patients:
        for side in ("left", "right"):
            d = IMG_DIR / pid / side
            if not d.is_dir():
                continue
            for f in sorted(d.iterdir()):
                if f.suffix.lower() in (".jpg", ".jpeg", ".png"):
                    paths.append(f)
                    owners.append(pid)
    return paths, owners


def build_dataset(label_source: str = "ted"):
    """(영상 경로, 소속 환자, 영상별 라벨, 환자 목록, 환자별 라벨) 을 만든다."""
    labels = load_labels_ted() if label_source == "ted" else load_labels_union()
    patients = [p for p in sorted(labels) if (IMG_DIR / p).is_dir()]
    paths, owners = list_images(patients)
    y_img = np.array([[labels[o][c] for c in CAS_ITEMS] for o in owners], dtype=np.float32)
    y_pat = np.array([[labels[p][c] for c in CAS_ITEMS] for p in patients], dtype=int)
    return paths, owners, y_img, patients, y_pat


def patient_folds(patients: list[str], y_pat: np.ndarray, n_splits: int = 5):
    """환자 단위 다중라벨 층화 K-겹. 시드를 고정해 모든 백본이 같은 분할을 쓴다."""
    from iterstrat.ml_stratifiers import MultilabelStratifiedKFold
    mskf = MultilabelStratifiedKFold(n_splits=n_splits, shuffle=True, random_state=SEED)
    idx = np.arange(len(patients))
    return [(tr, va) for tr, va in mskf.split(idx, y_pat)]


# ---------------------------------------------------------------- 백본
BACKBONES = {
    # 이름: (종류, 식별자, 입력 해상도, 특징 차원)
    "resnet18":  ("tv",   "resnet18",                             224, 512),
    "resnet50":  ("tv",   "resnet50",                             224, 2048),
    "dinov2_s":  ("timm", "vit_small_patch14_reg4_dinov2.lvd142m", 518, 384),
    "dinov2_b":  ("timm", "vit_base_patch14_reg4_dinov2.lvd142m",  518, 768),
}

_MEAN = (0.485, 0.456, 0.406)
_STD = (0.229, 0.224, 0.225)


def _preprocess(path: Path, size: int) -> torch.Tensor:
    """중앙 정사각 크롭 후 리사이즈. 안구가 화면 중앙에 있으므로 종횡비를 왜곡하지 않는다."""
    im = Image.open(path).convert("RGB")
    w, h = im.size
    s = min(w, h)
    im = im.crop(((w - s) // 2, (h - s) // 2, (w - s) // 2 + s, (h - s) // 2 + s))
    im = im.resize((size, size), Image.BICUBIC)
    x = torch.from_numpy(np.asarray(im, dtype=np.float32) / 255.0).permute(2, 0, 1)
    return (x - torch.tensor(_MEAN)[:, None, None]) / torch.tensor(_STD)[:, None, None]


def _build_backbone(name: str):
    kind, ident, size, dim = BACKBONES[name]
    if kind == "tv":
        import torchvision.models as tvm
        weights = {"resnet18": tvm.ResNet18_Weights.IMAGENET1K_V1,
                   "resnet50": tvm.ResNet50_Weights.IMAGENET1K_V2}[ident]
        m = getattr(tvm, ident)(weights=weights)
        m.fc = nn.Identity()
    else:
        import timm
        m = timm.create_model(ident, pretrained=True, num_classes=0)
    return m.eval(), size, dim


@torch.no_grad()
def extract_features(name: str, paths: list[Path], verbose: bool = True) -> np.ndarray:
    """백본별 특징을 뽑아 캐시한다. 같은 (백본, 영상목록)이면 재계산하지 않는다."""
    key = hashlib.md5(("|".join(str(p) for p in paths)).encode()).hexdigest()[:10]
    cache_file = CACHE / f"{name}_{key}.npy"
    if cache_file.exists():
        if verbose:
            print(f"  [캐시] {cache_file.name}")
        return np.load(cache_file)

    model, size, dim = _build_backbone(name)
    torch.set_grad_enabled(False)
    feats = np.zeros((len(paths), dim), dtype=np.float32)
    for i, p in enumerate(paths):
        feats[i] = model(_preprocess(p, size).unsqueeze(0)).squeeze(0).numpy()
        if verbose and (i + 1) % 25 == 0:
            print(f"  {name}: {i+1}/{len(paths)}", flush=True)
    np.save(cache_file, feats)
    return feats


# ---------------------------------------------------------------- 헤드
class Head(nn.Module):
    """논문과 동일한 항목별 독립 이진 헤드. hidden=0이면 선형 프로브."""

    def __init__(self, dim: int, hidden: int = 256, dropout: float = 0.2):
        super().__init__()
        if hidden:
            self.net = nn.Sequential(
                nn.LayerNorm(dim), nn.Linear(dim, hidden), nn.GELU(),
                nn.Dropout(dropout), nn.Linear(hidden, 1))
        else:
            self.net = nn.Sequential(nn.LayerNorm(dim), nn.Linear(dim, 1))

    def forward(self, x):
        return self.net(x).squeeze(-1)


def train_head(xt, yt, dim, hidden=256, epochs=200, lr=1e-3, wd=1e-2, seed=SEED):
    """단일 항목 헤드를 캐시된 특징 위에서 학습한다."""
    torch.manual_seed(seed)
    head = Head(dim, hidden)
    pos, neg = float(yt.sum()), float(len(yt) - yt.sum())
    pw = torch.tensor(max(neg / max(pos, 1.0), 1.0))
    loss_fn = nn.BCEWithLogitsLoss(pos_weight=pw)
    opt = torch.optim.AdamW(head.parameters(), lr=lr, weight_decay=wd)
    xt_t, yt_t = torch.from_numpy(xt), torch.from_numpy(yt)
    head.train()
    for _ in range(epochs):
        opt.zero_grad()
        loss_fn(head(xt_t), yt_t).backward()
        opt.step()
    head.eval()
    return head


@torch.no_grad()
def predict(head, x):
    return torch.sigmoid(head(torch.from_numpy(x))).numpy()
